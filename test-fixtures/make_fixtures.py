#!/usr/bin/env python3
"""确定性生成合成测试素材（文档/音频/杂项 + Excel 补充样例）。

真原件（图片/PDF/视频/原始Excel）直接存在母本里；本脚本只负责
可程序化生成的部分，保证每次生成的内容一致、可随代码演进重生成。
用法：python3 make_fixtures.py [母本目录，默认脚本所在目录]
依赖：python-docx openpyxl Pillow；音频需 ffmpeg（找不到则跳过并提示）。
"""
import csv
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent

def main():
    docs = ROOT / "Test Docs"; docs.mkdir(exist_ok=True)
    misc = ROOT / "Test Misc"; misc.mkdir(exist_ok=True)
    audio = ROOT / "Test Audio"; audio.mkdir(exist_ok=True)
    xls = ROOT / "Test Excel"

    # --- 文档 ---
    (docs / "销售说明_gbk.txt").write_bytes("这是一份需要转码的说明文档，包含中文内容。".encode("gbk"))
    (docs / "长日志.txt").write_text(
        "\n".join(f"log line {i}" for i in range(1, 3501)), encoding="utf-8")
    (docs / "待清理.txt").write_text("第一行   \n\n\n\n第二行\t\n第三行  \n", encoding="utf-8")

    import docx
    for name, title in [("采购合同A.docx", "采购合同"), ("采购合同B.docx", "服务合同")]:
        d = docx.Document()
        d.add_paragraph(title)
        d.add_paragraph("甲方：星辰科技")
        d.add_paragraph(f"乙方：供应商{name[4]}")
        d.save(docs / name)
    tpl = docx.Document()
    tpl.add_paragraph("聘书")
    tpl.add_paragraph("兹聘请 {{姓名}} 担任 {{岗位}} 一职。")
    tpl.save(docs / "聘书模板.docx")

    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["姓名", "岗位"])
    ws.append(["张三", "产品经理"])
    ws.append(["李四", "工程师"])
    wb.save(docs / "名单.xlsx")

    # --- Excel 补充样例（历史上曾有、后来丢失的两类）---
    rows = list(csv.reader(open(xls / "file_example_XLS_100.csv",
                                encoding="utf-8", errors="ignore")))
    wb = Workbook(); ws = wb.active
    for r in rows:
        ws.append(r[:8])
    wb.save(xls / "样例数据.xlsx")

    wb = Workbook(); ws = wb.active
    ws.append(["姓名-电话"])
    ws.append(["张三-13800001111"])
    ws.append(["李四-13900002222"])
    wb.save(xls / "待拆分示例.xlsx")

    wb = Workbook(); ws = wb.active
    ws.append(["分公司", "联系方式"])
    ws.append(["北京", "联系张三 13812345678 或 zhang@corp.com"])
    ws.append(["上海", "li@corp.com；备用 13998765432"])
    wb.save(xls / "待提取示例.xlsx")

    # --- 杂项 ---
    from PIL import Image, ImageDraw
    st = Image.new("RGBA", (300, 300), (0, 0, 0, 0))
    d = ImageDraw.Draw(st)
    d.ellipse([10, 10, 290, 290], outline=(200, 30, 30, 230), width=14)
    d.text((110, 135), "印章", fill=(200, 30, 30, 230))
    st.save(misc / "印章.png")

    # 多行 payload 二维码（vCard 名片，W-48）：payload 与 validator 同源。
    try:
        import qrcode
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        from qa_structural_cases import MULTILINE_QR_FIXTURE, MULTILINE_QR_PAYLOAD
        qrcode.make(MULTILINE_QR_PAYLOAD).save(misc / MULTILINE_QR_FIXTURE)
    except ImportError:
        print("提示：未找到 qrcode 库，跳过多行二维码素材（可用已提交的预生成 PNG）")

    vid = ROOT / "Test Video"
    (vid / "file_example_MP4_1280_10MG.srt").write_text(
        "1\n00:00:00,000 --> 00:00:01,000\n测试字幕第一行\n\n"
        "2\n00:00:01,200 --> 00:00:02,000\nSecond subtitle line\n", encoding="utf-8")

    # --- 音频（正弦波，确定性）---
    ffmpeg = shutil.which("ffmpeg") or str(
        Path.home() / "Library/Application Support/FinderSidebar/bin/ffmpeg")
    if Path(ffmpeg).exists():
        for name, freq, dur in [("clip_a.wav", 440, 4), ("clip_b.wav", 660, 3)]:
            subprocess.run([ffmpeg, "-hide_banner", "-loglevel", "error",
                            "-f", "lavfi", "-i", f"sine=frequency={freq}:duration={dur}",
                            "-ar", "22050", str(audio / name), "-y"], check=True)
    else:
        print("提示：未找到 ffmpeg，跳过音频素材（不影响其他部分）")

    print(f"合成素材生成完成 → {ROOT}")

if __name__ == "__main__":
    main()
