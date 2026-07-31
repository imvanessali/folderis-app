"""Semantic assertions for the established 47 Folderis real-fixture cases."""
from __future__ import annotations

import csv
import hashlib
import json
import subprocess
import wave
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path


def _first(outputs: list[Path], suffix: str | None = None) -> Path:
    matches = [path for path in outputs if suffix is None or path.suffix.lower() == suffix]
    if not matches:
        raise AssertionError(f"no new {suffix or ''} output")
    return matches[0]


def _image(path: Path):
    from PIL import Image
    return Image.open(path)


def _pdf(path: Path, password: str | None = None):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if password and reader.is_encrypted:
        reader.decrypt(password)
    return reader


def _workbook(path: Path, read_only: bool = True):
    import openpyxl
    return openpyxl.load_workbook(path, read_only=read_only, data_only=False)


def _media(ffprobe: Path, path: Path) -> dict:
    result = subprocess.run(
        [str(ffprobe), "-v", "error", "-show_format", "-show_streams", "-of", "json", str(path)],
        capture_output=True,
        text=True,
    )
    if result.returncode:
        raise AssertionError(result.stderr[-400:])
    return json.loads(result.stdout)


def _duration(probe: dict) -> float:
    return float(probe["format"]["duration"])


def _sha(path: Path, algorithm: str) -> str:
    digest = hashlib.new(algorithm)
    digest.update(path.read_bytes())
    return digest.hexdigest()


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    return "".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))


def validate_structural_case(
    case: dict,
    *,
    work: Path,
    cwd: Path,
    inputs: list[Path],
    outputs: list[Path],
    stdout: str,
    ffprobe: Path,
    run_skill,
) -> dict:
    """Return command-specific evidence or raise AssertionError."""
    id_ = case["id"]
    evidence: dict = {}

    if id_ == "W-01":
        out = _first(outputs)
        with _image(inputs[0]) as source, _image(out) as result:
            assert result.size == source.size and result.tobytes() != source.convert(result.mode).tobytes()
            evidence = {"format": result.format, "dimensions": list(result.size), "content_changed": True}
    elif id_ == "W-02":
        out = _first(outputs)
        with _image(inputs[0]) as source, _image(out) as result:
            assert result.size == (source.height, source.width)
            evidence = {"source_dimensions": list(source.size), "output_dimensions": list(result.size)}
    elif id_ == "W-03":
        out = _first(outputs)
        with _image(out) as result:
            assert result.format == "JPEG" and out.stat().st_size <= inputs[0].stat().st_size
            evidence = {"format": result.format, "source_bytes": inputs[0].stat().st_size, "output_bytes": out.stat().st_size}
    elif id_ == "W-04":
        out = _first(outputs, ".png")
        with _image(out) as result:
            assert result.format == "PNG"
            evidence = {"format": result.format, "dimensions": list(result.size)}
    elif id_ == "W-05":
        out = _first(outputs)
        with _image(out) as result:
            assert result.width == 1920
            evidence = {"width": result.width, "height": result.height}
    elif id_ == "W-06":
        assert stdout.strip()
        evidence = {"ocr_output_characters": len(stdout.strip())}
    elif id_ == "W-07":
        out = _first(outputs)
        with _image(inputs[0]) as source, _image(out) as result:
            assert result.width == source.width and result.height == source.height * 2
            evidence = {"dimensions": list(result.size), "direction": "vertical"}
    elif id_ == "W-08":
        out = _first(outputs)
        with _image(out) as result:
            assert result.size == (800, 600)
            evidence = {"dimensions": list(result.size), "origin": [100, 50]}
    elif id_ == "W-09":
        assert len(outputs) == 3
        with _image(inputs[0]) as source:
            total = 0
            widths = set()
            for output in outputs:
                with _image(output) as part:
                    total += part.height
                    widths.add(part.width)
            assert widths == {source.width} and total == source.height
            evidence = {"part_count": len(outputs), "total_height": total}
    elif id_ == "W-10":
        out = _first(outputs, ".ico")
        with _image(out) as result:
            assert result.format == "ICO"
            evidence = {"format": result.format, "dimensions": list(result.size)}
    elif id_ == "W-11":
        out = _first(outputs, ".docx")
        merged_text = _docx_text(out)
        source_texts = [_docx_text(path).strip() for path in inputs]
        assert all(not text or text in merged_text for text in source_texts)
        evidence = {"format": "DOCX", "input_documents_found": len(source_texts)}
    elif id_ == "W-12":
        out = _first(outputs, ".pdf")
        pages = len(_pdf(out).pages)
        assert out.read_bytes().startswith(b"%PDF") and pages > 0
        evidence = {"format": "PDF", "pages": pages}
    elif id_ == "W-13":
        out = _first(outputs, ".txt")
        text = out.read_text(encoding="utf-8")
        assert all(line == line.rstrip() for line in text.splitlines()) and "\n\n\n" not in text
        evidence = {"trimmed": True, "collapsed_blank_lines": True, "characters": len(text)}
    elif id_ == "W-14":
        out = _first(outputs, ".txt")
        text = out.read_text(encoding="utf-8")
        assert text and "\ufffd" not in text
        evidence = {"encoding": "utf-8", "characters": len(text)}
    elif id_ == "W-15":
        assert len(outputs) == 4
        line_counts = [len(path.read_text(encoding="utf-8").splitlines()) for path in outputs]
        assert all(0 < count <= 1000 for count in line_counts) and sum(line_counts) == 3500
        evidence = {"part_count": 4, "line_counts": line_counts}
    elif id_ == "W-16":
        data = _workbook(inputs[1])
        expected = data.active.max_row - 1
        assert len(outputs) == expected and all(path.suffix.lower() == ".docx" for path in outputs)
        evidence = {"data_rows": expected, "generated_documents": len(outputs)}
    elif id_ == "W-17":
        out = _first(outputs, ".pdf")
        pages = len(_pdf(out).pages)
        assert pages == 8
        evidence = {"format": "PDF", "pages": pages}
    elif id_ == "W-18":
        out = _first(outputs, ".pdf")
        source_pages, output_pages = len(_pdf(inputs[0]).pages), len(_pdf(out).pages)
        assert output_pages == source_pages
        evidence = {"source_pages": source_pages, "output_pages": output_pages, "output_bytes": out.stat().st_size}
    elif id_ == "W-19":
        out = _first(outputs, ".pdf")
        reader = _pdf(out)
        assert reader.is_encrypted and reader.decrypt("abc123")
        evidence = {"encrypted": True, "password_verified": True, "pages": len(reader.pages)}
    elif id_ == "W-20":
        out = _first(outputs, ".pdf")
        reader = _pdf(out)
        assert not reader.is_encrypted and len(reader.pages) == 4
        evidence = {"encrypted": False, "pages": len(reader.pages)}
    elif id_ == "W-21":
        assert len(outputs) == 4 and all(len(_pdf(path).pages) == 1 for path in outputs)
        evidence = {"part_count": 4, "pages_per_part": [len(_pdf(path).pages) for path in outputs]}
    elif id_ == "W-22":
        out = _first(outputs, ".pdf")
        source = _pdf(inputs[0])
        result = _pdf(out)
        assert len(result.pages) == 3
        source_text = [source.pages[index - 1].extract_text() for index in [3, 1, 2]]
        result_text = [page.extract_text() for page in result.pages]
        assert result_text == source_text
        evidence = {"pages": 3, "verified_order": [3, 1, 2]}
    elif id_ == "W-23":
        out = _first(outputs, ".txt")
        text = out.read_text(encoding="utf-8")
        assert len(text.strip()) > 100 and "第 1 页" in text
        evidence = {"format": "TXT", "characters": len(text)}
    elif id_ == "W-24":
        pngs = [path for path in outputs if path.suffix.lower() == ".png"]
        assert len(pngs) == 4
        dimensions = []
        for path in pngs:
            with _image(path) as result:
                assert result.format == "PNG"
                dimensions.append(list(result.size))
        evidence = {"format": "PNG", "page_images": len(pngs), "dimensions": dimensions}
    elif id_ == "W-25":
        out = _first(outputs, ".pdf")
        assert len(_pdf(out).pages) == 4 and _sha(out, "sha256") != _sha(inputs[0], "sha256")
        evidence = {"pages": 4, "content_changed": True, "text": "机密"}
    elif id_ == "W-26":
        out = _first(outputs, ".pdf")
        assert len(_pdf(out).pages) == 4 and _sha(out, "sha256") != _sha(inputs[0], "sha256")
        evidence = {"pages": 4, "content_changed": True, "page": "last", "position": "br"}
    elif id_ == "W-27":
        out = _first(outputs, ".csv")
        with out.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        assert len(rows) == 101 and len(rows[0]) > 1
        evidence = {"format": "CSV/UTF-8-SIG", "rows": len(rows), "columns": len(rows[0])}
    elif id_ == "W-28":
        source = _workbook(inputs[0])
        xlsx = [path for path in outputs if path.suffix.lower() == ".xlsx"]
        assert len(xlsx) == len(source.sheetnames)
        evidence = {"source_sheets": source.sheetnames, "output_workbooks": len(xlsx)}
    elif id_ == "W-29":
        out = _first(outputs, ".xlsx")
        source, result = _workbook(inputs[0]), _workbook(out)
        assert result.active.max_row <= source.active.max_row and result.active.max_column == source.active.max_column
        evidence = {"source_rows": source.active.max_row, "output_rows": result.active.max_row, "columns": result.active.max_column}
    elif id_ == "W-30":
        assert "仅在 " in stdout and "键相同但内容不同" in stdout
        evidence = {"comparison_output_characters": len(stdout.strip())}
    elif id_ == "W-31":
        out = _first(outputs, ".xlsx")
        source, result = _workbook(inputs[0]), _workbook(out)
        assert result.active.max_column > source.active.max_column
        evidence = {"source_columns": source.active.max_column, "output_columns": result.active.max_column}
    elif id_ == "W-32":
        out = _first(outputs, ".xlsx")
        result = _workbook(out)
        assert result.active.max_row > 1 and result.active.max_column >= 2
        evidence = {"rows": result.active.max_row, "columns": result.active.max_column, "aggregation": "sum"}
    elif id_ == "W-33":
        out = _first(outputs, ".xlsx")
        result, target = _workbook(out), _workbook(inputs[1])
        assert result.active.max_row == target.active.max_row
        evidence = {"rows": result.active.max_row, "key_column": 1}
    elif id_ == "W-34":
        assert "共 " in stdout and " 行" in stdout and " 列" in stdout
        evidence = {"preview_output": stdout.strip()[:300]}
    elif id_ == "W-35":
        xlsx = [path for path in outputs if path.suffix.lower() == ".xlsx"]
        assert len(xlsx) > 1 and all(_workbook(path).active.max_row > 1 for path in xlsx)
        evidence = {"output_workbooks": len(xlsx), "by_column": 5}
    elif id_ == "W-36":
        out = _first(outputs, ".xlsx")
        source, result = _workbook(inputs[0]), _workbook(out)
        assert result.active.max_row == source.active.max_row and result.active.max_column >= source.active.max_column
        evidence = {"rows": result.active.max_row, "columns": result.active.max_column}
    elif id_ == "W-37":
        out = _first(outputs, ".xlsx")
        result = _workbook(out)
        assert result.active.max_row > 1
        evidence = {"extracted_rows": result.active.max_row - 1, "pattern": "email"}
    elif id_ == "W-38":
        out = _first(outputs, ".mkv")
        probe = _media(ffprobe, out)
        assert "matroska" in probe["format"]["format_name"]
        evidence = {"format": probe["format"]["format_name"], "duration_seconds": _duration(probe)}
    elif id_ == "W-39":
        out = _first(outputs, ".mp3")
        probe = _media(ffprobe, out)
        audio = next(stream for stream in probe["streams"] if stream["codec_type"] == "audio")
        assert audio["codec_name"] == "mp3"
        evidence = {"codec": audio["codec_name"], "duration_seconds": _duration(probe)}
    elif id_ == "W-40":
        frames = [path for path in outputs if path.suffix.lower() in {".jpg", ".jpeg"}]
        assert len(frames) >= 3
        evidence = {"frame_count": len(frames), "every_seconds": 10}
    elif id_ == "W-41":
        out = _first(outputs, ".gif")
        with _image(out) as result:
            frame_count = getattr(result, "n_frames", 1)
            milliseconds = 0
            for index in range(frame_count):
                result.seek(index)
                milliseconds += result.info.get("duration", 0)
            assert result.format == "GIF" and result.width == 480 and 4000 <= milliseconds <= 5500
            evidence = {"format": "GIF", "dimensions": [result.width, result.height], "frames": frame_count, "duration_seconds": milliseconds / 1000}
    elif id_ == "W-42":
        out = _first(outputs, ".m4a")
        probe = _media(ffprobe, out)
        assert any(stream["codec_type"] == "audio" for stream in probe["streams"])
        evidence = {"format": probe["format"]["format_name"], "duration_seconds": _duration(probe)}
    elif id_ == "W-43":
        out = _first(outputs, ".wav")
        with wave.open(str(out)) as result:
            actual = result.getnframes() / result.getframerate()
            channels, rate = result.getnchannels(), result.getframerate()
        assert abs(actual - 7.0) <= 0.1
        evidence = {"format": "WAV", "duration_seconds": actual, "channels": channels, "sample_rate": rate}
    elif id_ == "W-44":
        out = _first(outputs, ".xlsx")
        workbook = _workbook(out)
        headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
        assert workbook.active.max_row > 1 and {"名称", "相对路径", "类型", "大小(字节)", "修改时间"} <= set(headers)
        evidence = {"format": "XLSX", "inventory_rows": workbook.active.max_row - 1, "headers": headers}
    elif id_ == "W-45":
        expected = {str(path): _sha(path, "md5") for path in inputs}
        assert all(digest in stdout for digest in expected.values()) and not outputs
        evidence = {"algorithm": "md5", "digests": expected}
    elif id_ == "W-46":
        out = _first(outputs, ".zip")
        assert zipfile.is_zipfile(out)
        with zipfile.ZipFile(out) as archive:
            infos = archive.infolist()
            assert infos and all(info.flag_bits & 0x1 for info in infos)
            archive.setpassword(b"abc123")
            assert archive.read(infos[0])
        evidence = {"format": "ZIP", "encrypted": True, "entries": [info.filename for info in infos]}
    elif id_ == "W-47":
        out = _first(outputs, ".png")
        with _image(out) as result:
            assert result.format == "PNG" and result.size == (512, 512)
        readback = run_skill(cwd, ["qr-read", str(out)])
        assert readback.returncode == 0 and "https://folderis.app" in readback.stdout
        evidence = {"format": "PNG", "dimensions": [512, 512], "decoded_content": "https://folderis.app"}
    else:
        raise AssertionError(f"missing semantic validator for {id_}")
    return evidence
