import argparse
import importlib.util
import tempfile
import unittest
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest import mock


ROOT = Path(__file__).resolve().parents[1]
SKILLS = ROOT / "skills" / "skills.py"


def load_skills():
    spec = importlib.util.spec_from_file_location("folderis_skills_output_tests", SKILLS)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class OutputVersioningTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.skills = load_skills()

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp_dir.name)
        self.source = self.folder / "movie.mp4"
        self.source.write_bytes(b"original-video")

    def tearDown(self):
        self.temp_dir.cleanup()

    def trim(self, output=None):
        args = argparse.Namespace(
            files=[str(self.source)],
            start="0",
            end="60",
            output=str(output) if output else None,
            fast=True,
        )

        def fake_run(command):
            target = Path(command[-1])
            target.write_bytes(f"generated:{target.name}".encode())

        with mock.patch.object(self.skills, "ffmpeg_or_die", return_value="ffmpeg"), \
             mock.patch.object(self.skills, "run", side_effect=fake_run):
            self.skills.cmd_av_trim(args)

    def test_repeated_trim_never_changes_earlier_outputs(self):
        self.trim()
        first = self.folder / "movie_clip.mp4"
        first_bytes = first.read_bytes()

        self.trim()
        second = self.folder / "movie_clip_2.mp4"
        second_bytes = second.read_bytes()

        self.trim()

        self.assertEqual(self.source.read_bytes(), b"original-video")
        self.assertEqual(first.read_bytes(), first_bytes)
        self.assertEqual(second.read_bytes(), second_bytes)
        self.assertTrue((self.folder / "movie_clip_3.mp4").exists())

    def test_occupied_explicit_output_is_versioned(self):
        requested = self.folder / "chosen.mp4"
        requested.write_bytes(b"keep-this-result")

        self.trim(requested)

        self.assertEqual(requested.read_bytes(), b"keep-this-result")
        self.assertTrue((self.folder / "chosen_2.mp4").exists())

    def test_repeated_text_processing_preserves_every_result(self):
        source = self.folder / "notes.txt"
        source.write_text("line 1  \n\n\nline 2\n", encoding="utf-8")
        args = argparse.Namespace(
            files=[str(source)],
            trim=True,
            collapse_blank=True,
            strip_all_blank=False,
            in_place=False,
        )

        self.skills.cmd_txt_clean(args)
        first = self.folder / "notes_clean.txt"
        first_bytes = first.read_bytes()
        self.skills.cmd_txt_clean(args)

        self.assertEqual(source.read_text(encoding="utf-8"), "line 1  \n\n\nline 2\n")
        self.assertEqual(first.read_bytes(), first_bytes)
        self.assertTrue((self.folder / "notes_clean_2.txt").exists())

    def test_split_archive_family_avoids_existing_parts(self):
        first_part = self.folder / "archive.7z.001"
        first_part.write_bytes(b"old-part")

        output = self.skills.available_family_output(self.folder / "archive.7z")

        self.assertEqual(output.name, "archive_2.7z")
        self.assertEqual(first_part.read_bytes(), b"old-part")

    def test_concurrent_name_reservations_are_unique(self):
        requested = self.folder / "result.txt"

        with ThreadPoolExecutor(max_workers=8) as pool:
            outputs = list(pool.map(
                lambda _: self.skills.available_output(requested), range(8)))

        self.assertEqual(len({str(path) for path in outputs}), 8)
        self.assertEqual(
            {path.name for path in outputs},
            {"result.txt", *{f"result_{index}.txt" for index in range(2, 9)}},
        )

    def test_case_and_unicode_equivalent_paths_share_reservations(self):
        upper = self.skills.available_output(self.folder / "Result.txt")
        lower = self.skills.available_output(self.folder / "result.txt")
        self.assertEqual(upper.name, "Result.txt")
        self.assertEqual(lower.name, "result_2.txt")

        nfc = "caf\u00e9.txt"
        nfd = unicodedata.normalize("NFD", nfc)
        first = self.skills.available_output(self.folder / nfc)
        second = self.skills.available_output(self.folder / nfd)
        self.assertEqual(first.name, nfc)
        self.assertTrue(second.name.endswith("_2.txt"))

    def test_split_archive_reserves_future_member_namespace(self):
        base = self.skills.available_family_output(self.folder / "archive.zip")
        member = self.skills.available_output(self.folder / "archive.z01")

        self.assertEqual(base.name, "archive.zip")
        self.assertEqual(member.name, "archive_2.z01")

    def test_docx_fill_versions_duplicate_names_within_one_batch(self):
        from docx import Document
        from openpyxl import Workbook

        template = self.folder / "template.docx"
        document = Document()
        document.add_paragraph("Hello {{value}}")
        document.save(template)

        data = self.folder / "data.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["name", "value"])
        sheet.append(["same", "first"])
        sheet.append(["same", "second"])
        workbook.save(data)

        args = argparse.Namespace(
            template=str(template),
            data=str(data),
            name_column="name",
            output_dir=None,
        )
        self.skills.cmd_docx_fill(args)

        output_dir = self.folder / "template_filled"
        first = output_dir / "template_same.docx"
        second = output_dir / "template_same_2.docx"
        self.assertTrue(first.exists())
        self.assertTrue(second.exists())
        self.assertIn("first", "\n".join(p.text for p in Document(first).paragraphs))
        self.assertIn("second", "\n".join(p.text for p in Document(second).paragraphs))


if __name__ == "__main__":
    unittest.main()
