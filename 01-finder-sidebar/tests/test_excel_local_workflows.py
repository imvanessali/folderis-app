import hashlib
import csv
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = ROOT.parent
SKILLS = ROOT / "skills" / "skills.py"
XLS_FIXTURE = PROJECT_ROOT / "test-fixtures" / "Test Excel" / "file_example_XLS_100.xls"
RESULT_MARKER = "__FOLDERIS_LOCAL_RESULT__"


class ExcelLocalWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def run_skill(self, *arguments):
        result = subprocess.run(
            [sys.executable, str(SKILLS), *map(str, arguments)],
            cwd=self.folder,
            text=True,
            capture_output=True,
            check=True,
        )
        marker_payload = result.stdout.rsplit(RESULT_MARKER, 1)[-1].strip()
        return json.loads(marker_payload)

    def test_xls_date_to_month_and_country_sheets_preserves_source(self):
        source = self.folder / "source.xls"
        shutil.copy2(XLS_FIXTURE, source)
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        request = "根据 Date 新增一个 Month 列，例如 July，按地区把数据拆到不同的 sheet 里"

        payload = self.run_skill(
            "local-dispatch", "--request", request, "--files", source)

        self.assertTrue(payload["matched"])
        self.assertTrue(payload["success"])
        self.assertEqual(payload["skill"], "excel-derive-split-sheets")
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)
        self.assertEqual(len(payload["outputs"]), 1)

        output = Path(payload["outputs"][0])
        workbook = openpyxl.load_workbook(output, data_only=True)
        self.assertEqual(set(workbook.sheetnames), {"United States", "Great Britain", "France"})
        self.assertEqual(sum(sheet.max_row - 1 for sheet in workbook.worksheets), 100)
        months = set()
        for sheet in workbook.worksheets:
            header = [cell.value for cell in sheet[1]]
            self.assertEqual(header.count("Month"), 1)
            month_column = header.index("Month") + 1
            months.update(
                sheet.cell(row=row, column=month_column).value
                for row in range(2, sheet.max_row + 1)
            )
        self.assertIn("October", months)

        first_bytes = output.read_bytes()
        second = self.run_skill(
            "local-dispatch", "--request", request, "--files", source)
        self.assertEqual(len(second["outputs"]), 1)
        self.assertTrue(Path(second["outputs"][0]).name.endswith("_2.xlsx"))
        self.assertEqual(output.read_bytes(), first_bytes)

    def test_dummy_inventory_is_unicode_safe_and_leaves_no_scripts(self):
        payload = self.run_skill(
            "local-dispatch", "--request", "生成 12 行模拟库存 Excel 测试数据", "--files")

        self.assertTrue(payload["matched"])
        self.assertTrue(payload["success"])
        self.assertEqual(payload["skill"], "excel-dummy")
        self.assertEqual(len(payload["outputs"]), 1)
        output = Path(payload["outputs"][0])
        self.assertEqual(output.name, "模拟库存数据.xlsx")
        workbook = openpyxl.load_workbook(output, data_only=True)
        self.assertEqual(workbook.active.title, "库存数据")
        self.assertEqual(workbook.active["B2"].value, "测试产品 1")
        self.assertEqual(workbook.active.max_row, 13)
        leftovers = [
            path for path in self.folder.iterdir()
            if path.suffix.lower() in {".py", ".ps1", ".vbs"}
        ]
        self.assertEqual(leftovers, [])

    def test_question_about_dummy_excel_does_not_create_a_file(self):
        payload = self.run_skill(
            "local-dispatch", "--request", "什么是模拟库存 Excel 测试数据？", "--files")

        self.assertFalse(payload["matched"])
        self.assertEqual(list(self.folder.iterdir()), [])

    def test_image_conversion_dispatch_creates_one_result(self):
        source = self.folder / "image.jpg"
        Image.new("RGB", (20, 10), "red").save(source)
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

        payload = self.run_skill(
            "local-dispatch", "--request", "转成 png", "--files", source)

        self.assertTrue(payload["matched"])
        self.assertTrue(payload["success"])
        self.assertEqual(payload["skill"], "img-convert")
        self.assertEqual(len(payload["outputs"]), 1)
        self.assertEqual([path.name for path in self.folder.glob("*.png")], ["image_converted.png"])
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)

    def test_csv_column_split_replaces_source_with_named_columns(self):
        source = self.folder / "contacts.csv"
        source.write_text(
            "联系人-电话,城市,备注\n张三-13800138000,北京,重点客户\n李四-13900139000,上海,\n",
            encoding="utf-8-sig",
        )
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

        subprocess.run(
            [sys.executable, str(SKILLS), "excel-transform", source.name,
             "--split-column", "1", "--delimiter", "-", "--new-names", "姓名,电话"],
            cwd=self.folder,
            text=True,
            capture_output=True,
            check=True,
        )

        output = self.folder / "contacts_transformed.csv"
        with output.open(newline="", encoding="utf-8-sig") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[0], ["姓名", "电话", "城市", "备注"])
        self.assertEqual(rows[1], ["张三", "13800138000", "北京", "重点客户"])
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)

    def test_csv_column_split_can_keep_source(self):
        source = self.folder / "contacts.csv"
        source.write_text("联系人-电话,城市\n张三-13800138000,北京\n", encoding="utf-8-sig")
        subprocess.run(
            [sys.executable, str(SKILLS), "excel-transform", source,
             "--split-column", "1", "--delimiter", "-",
             "--new-names", "姓名,电话", "--keep-source"],
            cwd=self.folder, text=True, capture_output=True, check=True)
        with (self.folder / "contacts_transformed.csv").open(
                newline="", encoding="utf-8-sig") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[0], ["联系人-电话", "姓名", "电话", "城市"])
        self.assertEqual(rows[1], ["张三-13800138000", "张三", "13800138000", "北京"])


if __name__ == "__main__":
    unittest.main()
