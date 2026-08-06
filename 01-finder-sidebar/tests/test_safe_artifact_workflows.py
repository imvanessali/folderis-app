import argparse
import hashlib
import importlib.util
import json
import subprocess
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock


ROOT = Path(__file__).resolve().parents[2]
SKILLS_PATH = ROOT / "01-finder-sidebar" / "skills" / "skills.py"
SPEC = importlib.util.spec_from_file_location("folderis_safe_artifact_skills", SKILLS_PATH)
skills = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(skills)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class SafeArtifactWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_multi_image_conversion_uses_versioned_output_and_uniform_names(self):
        from PIL import Image
        sources = []
        for name, color, fmt in [("alpha.jpg", "red", "JPEG"), ("beta.png", "blue", "PNG")]:
            path = self.folder / name
            Image.new("RGB", (24, 18), color).save(path, fmt)
            sources.append(path)
        originals = {path.name: digest(path) for path in sources}
        args = argparse.Namespace(files=[str(path) for path in sources], to="webp", output=None,
                                  overwrite_original=False)

        skills.cmd_img_convert(args)
        skills.cmd_img_convert(args)

        self.assertEqual(sorted(path.name for path in (self.folder / "Output").iterdir()),
                         ["alpha_converted.webp", "beta_converted.webp"])
        self.assertEqual(sorted(path.name for path in (self.folder / "Output_2").iterdir()),
                         ["alpha_converted.webp", "beta_converted.webp"])
        self.assertEqual({path.name: digest(path) for path in sources}, originals)

    def test_text_cleanup_removes_boundary_blanks_and_preserves_bom_inputs(self):
        first = self.folder / "one.txt"
        second = self.folder / "two.txt"
        dirty = "\n\nFirst   \n\n \nSecond\t\n\n"
        first.write_text(dirty, encoding="utf-8-sig")
        second.write_text(dirty.replace("First", "Alpha"), encoding="utf-8")
        originals = {path.name: digest(path) for path in (first, second)}
        args = argparse.Namespace(files=[str(first), str(second)], trim=True,
                                  collapse_blank=True, strip_all_blank=False,
                                  in_place=False, overwrite_original=False)

        skills.cmd_txt_clean(args)

        output = self.folder / "Output"
        self.assertEqual((output / "one_cleaned.txt").read_text(encoding="utf-8-sig"),
                         "First\n\nSecond\n")
        self.assertEqual((output / "two_cleaned.txt").read_text(encoding="utf-8"),
                         "Alpha\n\nSecond\n")
        self.assertEqual({path.name: digest(path) for path in (first, second)}, originals)

    def test_excel_diff_always_writes_a_detailed_workbook_and_versions_it(self):
        import openpyxl
        old = self.folder / "old.xlsx"
        new = self.folder / "new.xlsx"
        for path, rows in [
            (old, [["ID", "Name", "Price"], ["A", "One", 10], ["B", "Two", 20]]),
            (new, [["ID", "Name", "Price"], ["A", "One", 12], ["B", "Two", 20], ["C", "Three", 30]]),
        ]:
            book = openpyxl.Workbook()
            sheet = book.active
            for row in rows:
                sheet.append(row)
            book.save(path)
        originals = (digest(old), digest(new))
        args = argparse.Namespace(file_a=str(old), file_b=str(new), columns="1", output=None)

        skills.cmd_excel_diff(args)
        skills.cmd_excel_diff(args)

        report = self.folder / "old_vs_new_diff.xlsx"
        self.assertTrue(report.exists())
        self.assertTrue((self.folder / "old_vs_new_diff_2.xlsx").exists())
        book = openpyxl.load_workbook(report, data_only=True)
        self.assertEqual(book.sheetnames, ["汇总", "仅旧表", "仅新表", "内容变化"])
        changes = list(book["内容变化"].iter_rows(min_row=2, values_only=True))
        self.assertIn(("A", "Price", 10, 12), changes)
        self.assertEqual((digest(old), digest(new)), originals)

    def test_safe_dedupe_builds_output_without_touching_inputs(self):
        a = self.folder / "a.bin"
        b = self.folder / "b.bin"
        c = self.folder / "c.bin"
        a.write_bytes(b"same")
        b.write_bytes(b"same")
        c.write_bytes(b"different")
        originals = {path.name: digest(path) for path in (a, b, c)}

        skills.cmd_dedupe_safe(argparse.Namespace(files=[str(a), str(b), str(c)]))

        output = self.folder / "Output"
        copied = sorted(path.name for path in output.iterdir() if path.suffix == ".bin")
        self.assertEqual(copied, ["a.bin", "c.bin"])
        self.assertTrue((output / "去重报告.xlsx").exists())
        self.assertEqual({path.name: digest(path) for path in (a, b, c)}, originals)

    def test_capture_time_rename_copies_into_output(self):
        from PIL import Image
        later = self.folder / "later.jpg"
        earlier = self.folder / "earlier.jpg"
        for path, stamp in [(later, "2026:02:01 10:00:00"), (earlier, "2026:01:01 10:00:00")]:
            exif = Image.Exif()
            exif[36867] = stamp
            Image.new("RGB", (16, 16), "green").save(path, exif=exif)
        originals = {path.name: digest(path) for path in (later, earlier)}
        args = argparse.Namespace(files=[str(later), str(earlier)], template="vacation_{n}",
                                  regex=None, replace=None, ext=None, start=1, pad=3,
                                  sort="capture-time", dry_run=False, overwrite_original=False)

        skills.cmd_rename_batch(args)

        self.assertEqual(digest(self.folder / "Output" / "vacation_001.jpg"), digest(earlier))
        self.assertEqual(digest(self.folder / "Output" / "vacation_002.jpg"), digest(later))
        self.assertEqual({path.name: digest(path) for path in (later, earlier)}, originals)

    def test_qr_read_writes_single_workbook_even_when_some_images_fail(self):
        import openpyxl
        first = self.folder / "first.png"
        second = self.folder / "second.png"
        first.write_bytes(b"one")
        second.write_bytes(b"two")
        stdout = f"{first}\thttps://example.test\n{second}\t(未识别到二维码/条形码)\n"
        result = subprocess.CompletedProcess([], 0, stdout=stdout, stderr="")
        args = argparse.Namespace(files=[str(first), str(second)], output=None)

        with mock.patch.object(skills, "IS_WIN", False), \
             mock.patch.object(skills, "helper", return_value=["fake-qr"]), \
             mock.patch.object(skills.subprocess, "run", return_value=result):
            skills.cmd_qr_read(args)

        report = self.folder / "二维码识别结果.xlsx"
        rows = list(openpyxl.load_workbook(report, data_only=True).active.values)
        self.assertEqual(rows[1], ("first.png", "https://example.test", "已识别"))
        self.assertEqual(rows[2], ("second.png", None, "未识别到二维码/条形码"))

    def test_precise_background_removal_preserves_sources_and_uses_output(self):
        sources = []
        for name, body in [("one.jpg", b"first"), ("two.png", b"second")]:
            path = self.folder / name
            path.write_bytes(body)
            sources.append(path)
        originals = {path.name: digest(path) for path in sources}
        calls = []

        def fake_remove(body):
            calls.append(body)
            return b"transparent-" + body

        args = argparse.Namespace(files=[str(path) for path in sources], engine="rembg")
        module = types.SimpleNamespace(remove=fake_remove)
        with mock.patch.dict("sys.modules", {"rembg": module}):
            skills.cmd_img_remove_bg(args)

        self.assertEqual(calls, [b"first", b"second"])
        self.assertEqual((self.folder / "Output" / "one_nobg.png").read_bytes(),
                         b"transparent-first")
        self.assertEqual((self.folder / "Output" / "two_nobg.png").read_bytes(),
                         b"transparent-second")
        self.assertEqual({path.name: digest(path) for path in sources}, originals)

    def test_material_independent_routes_cover_report_generating_cases(self):
        files = []
        for name in ["left.xlsx", "right.xlsx", "dirty-a.txt", "dirty-b.txt", "voice.wav", "code.png"]:
            path = self.folder / name
            path.write_bytes(b"fixture")
            files.append(path)
        self.assertEqual(skills._strict_local_plan("按ID列对比两个表的差异", files[:2])["skill"], "excel-diff")
        self.assertEqual(skills._strict_local_plan("清理行尾空格和连续空行", files[2:4])["skill"], "txt-clean")
        self.assertEqual(skills._strict_local_plan("将录音转写为文字", [files[4]])["skill"], "audio-transcribe")
        english = skills._strict_local_plan("转写这段英文录音", [files[4]])
        self.assertEqual(english["args"][-2:], ["--locale", "en-US"])
        self.assertEqual(skills._strict_local_plan("识别截图中的二维码", [files[5]])["skill"], "qr-read")
        self.assertIsNone(skills._strict_local_plan("将录音转写并总结为会议纪要", [files[4]]))

        first_doc = self.folder / "01.docx"
        second_doc = self.folder / "02.txt"
        first_doc.write_bytes(b"docx-placeholder")
        second_doc.write_text("text", encoding="utf-8")
        merge = skills._strict_local_plan("按文件名顺序合并成一份投资者资料包", [second_doc, first_doc])
        self.assertEqual(merge["skill"], "merge-docs")
        self.assertEqual(Path(merge["args"][-1]).name, "投资者资料包.docx")

    def test_local_dispatch_reports_files_nested_in_output(self):
        first = self.folder / "first.txt"
        second = self.folder / "second.txt"
        first.write_text("\n\nA  \n\n\nB\n", encoding="utf-8")
        second.write_text("\n\nC\t\n\n\nD\n", encoding="utf-8")
        completed = subprocess.run(
            ["python3", str(SKILLS_PATH), "local-dispatch", "--request",
             "清理行尾空格和连续空行", "--files", str(first), str(second)],
            cwd=self.folder, capture_output=True, text=True, check=True,
        )
        marker = "__FOLDERIS_LOCAL_RESULT__"
        payload = json.loads(completed.stdout.rsplit(marker, 1)[1])
        self.assertTrue(payload["matched"])
        self.assertTrue(payload["success"])
        self.assertEqual([Path(path).parent.name for path in payload["outputs"]],
                         ["Output", "Output"])
        self.assertEqual([Path(path).name for path in payload["outputs"]],
                         ["first_cleaned.txt", "second_cleaned.txt"])

    def test_zip_encrypt_uses_aes_256_safe_relative_paths_and_versions_output(self):
        try:
            import pyzipper
        except ImportError:
            self.skipTest("pyzipper is verified with the packaged Folderis runtime")
        first = self.folder / "合同.txt"
        first.write_text("secret", encoding="utf-8")
        nested = self.folder / "资料"
        nested.mkdir()
        (nested / "附件.txt").write_text("nested", encoding="utf-8")
        originals = {path: digest(path) for path in (first, nested / "附件.txt")}
        args = argparse.Namespace(files=[str(first.resolve()), str(nested.resolve())],
                                  password="abc123", output=None)

        skills.cmd_zip_encrypt(args)
        skills.cmd_zip_encrypt(args)

        archive_path = self.folder / "加密压缩包.zip"
        self.assertTrue((self.folder / "加密压缩包_2.zip").exists())
        with pyzipper.AESZipFile(archive_path) as archive:
            archive.setpassword(b"abc123")
            names = archive.namelist()
            self.assertEqual(names, ["合同.txt", "资料/附件.txt"])
            self.assertEqual(archive.read("合同.txt"), b"secret")
            self.assertTrue(all(info.flag_bits & 0x1 for info in archive.infolist()))
        self.assertTrue(all(not Path(name).is_absolute() and ".." not in Path(name).parts
                            for name in names))
        self.assertEqual({path: digest(path) for path in originals}, originals)


if __name__ == "__main__":
    unittest.main()
